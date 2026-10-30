#!/usr/bin/env python3
"""Write the Pro (native, C2 Proficiency) vocabulary and grammar summary as Word, same layout as build.py.

Run (repo root):  python3 tools/study-summary/pro.py OUT.docx data.json
data.json: from dump-pro.js (4 ESO LS2, read from data/native/*.js) or hand-written (pro-ls1-4eso.json, LS1 has no Pro pages).
Shape: {course, ls, extension?, lps: [{n, title, vocab: [{name, items: [[word, meaning, example]]}],
grammar: [{title, how: [html], practice: [[sentence, answer, why]]}]}]}. Convert to PDF with LibreOffice and stamp it with
tools/stamp-copyright.py, as for the other routes (build-pro.sh does all of it).
"""
import json, sys
from docx.shared import Cm, Pt

import build
from build import ACCENT, GREY, COURSE, html_paras, keep_with_next, run


class Pro(build.Builder):
    def pro(self):
        d = self.d
        self.first = False
        self.para("%s · %s" % (COURSE[self.course], d["ls"]), bold=True, size=9, color=ACCENT)
        p = self.para("Pro route · Vocabulary and grammar (C2 Proficiency)", bold=True, size=20, color=ACCENT)
        p.paragraph_format.space_after = Pt(6)
        self.para(d.get("intro") or "All the vocabulary and grammar of the Pro Learning Paths, with the explanations as they appear on the site.")
        self.box("How to study it", [
            [("First all the vocabulary, then all the grammar, Learning Path by Learning Path. " + d.get("order", ""), False)],
            [("Vocabulary: ", True), ("cover the meaning column and define each word in your own words; then write a new example "
                                       "sentence of your own, as precise as the one in the table.", False)],
            [("Grammar: ", True), ("read how it works, cover the answer column and do the transformations; then check and read "
                                    "why. In the exam, small spelling mistakes cost the mark.", False)],
        ])
        self.doc.add_heading("1 · Vocabulary", 1)
        for lp in d["lps"]:
            if not lp["vocab"]:
                continue
            keep_with_next(self.doc.add_heading("Learning Path %d · %s" % (lp["n"], lp["title"]), 2))
            rows, group_rows = [], []
            for g in lp["vocab"]:
                group_rows.append(len(rows))
                rows.append(["%s · %d" % (g["name"], len(g["items"]))])
                rows += [list(x) for x in g["items"]]
            self.table(("Word or expression", "Meaning", "Example"), rows, (0.27, 0.37, 0.36), group_rows, bold_col=0, italic_col=2)
        self.doc.add_heading("2 · Grammar", 1)
        if d.get("gintro"):
            self.para(d["gintro"], color=GREY)
        for lp in d["lps"]:
            if not lp["grammar"]:
                continue
            keep_with_next(self.doc.add_heading("Learning Path %d · %s" % (lp["n"], lp["title"]), 2))
            for b in lp["grammar"]:
                keep_with_next(self.doc.add_heading(b["title"], 3))
                keep_with_next(self.para("How it works", bold=True, size=10, color=ACCENT))
                for html in b["how"]:
                    for kind, runs in html_paras(html):
                        q = self.doc.add_paragraph(style="List Bullet")
                        q.paragraph_format.space_after = Pt(2)
                        for text, fmt in runs:
                            run(q, text, bold="bold" in fmt, italic="italic" in fmt)
                p = self.para("Practice from the Learning Path, with answers", bold=True, size=10, color=ACCENT)
                p.paragraph_format.space_before = Pt(6); keep_with_next(p)
                self.table(("Sentence", "Answer", "Why"), b["practice"], (0.52, 0.22, 0.26), bold_col=1)


def main():
    out, data = sys.argv[1], json.load(open(sys.argv[2], encoding="utf-8"))
    build.LS = data["ls"]  # footer
    b = Pro(data)
    b.pro()
    b.doc.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
