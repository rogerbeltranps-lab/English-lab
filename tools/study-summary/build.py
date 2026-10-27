#!/usr/bin/env python3
"""Write the LS1 vocabulary and grammar summary as Word, from the data of dump.js.

Run (repo root):  python3 tools/study-summary/build.py OUT.docx data.json [insecure|safe|daring]
Without a route it writes the teacher's Word (a map of what each route learns, then one part per route, each on its
own page); with a route it writes that route's student version (convert it to PDF with LibreOffice and stamp it with
tools/stamp-copyright.py). Same layout as the LS2 summaries. Needs python-docx.
"""
import datetime, json, re, sys
from html.parser import HTMLParser
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROUTES = ["insecure", "safe", "daring"]
NAME = {"insecure": "Insecure", "safe": "Safe", "daring": "Daring"}
ACCENT, LIGHT, BORDER, TEXT, GREY = "1E5B66", "EEF4F5", "B9C6CB", "1F2A30", "5F6B73"
COURSE = {"2eso": "2 ESO", "4eso": "4 ESO"}
LS = "LS1 My Real Routine"
COPYRIGHT = "© %d Roger Beltran · English Lab · All rights reserved" % datetime.date.today().year

T = {
    "en": dict(
        route="{r} route · Vocabulary and grammar",
        intro="All the vocabulary and grammar of the Learning Paths on the {r} route, with the explanations as they appear on the site.",
        box="How to study it",
        order="First all the vocabulary, then all the grammar, Learning Path by Learning Path. Learning Path 1 (a warm-up) and Learning Path 6 (the final product) have no new vocabulary or grammar.",
        vtip=("Vocabulary: ", "cover the Catalan column and say what each word means; then read the example aloud."),
        gtip=("Grammar: ", "read how it works, cover the answer column and complete the sentences; then check and read why."),
        vocab="1 · Vocabulary", grammar="2 · Grammar",
        gintro="Learning Path 1 is a warm-up and Learning Path 6 is the final product: they have no new grammar.",
        cols=("Word or expression", "In Catalan", "Example"), how="How it works",
        practice="Practice from the Learning Path, with answers", pcols=("Sentence", "Answer", "Why"),
    ),
    "ca": dict(
        route="Ruta {r} · Vocabulari i gramàtica",
        intro="Tot el vocabulari i la gramàtica dels Learning Paths de la ruta {r}, amb les explicacions tal com surten al web.",
        box="Com estudiar-ho",
        order="Primer tot el vocabulari i després tota la gramàtica, Learning Path per Learning Path. El Learning Path 1 (l’escalfament) i el 6 (el producte final) no tenen vocabulari ni gramàtica nous.",
        vtip=("Vocabulari: ", "tapa la columna del català i prova de dir què vol dir cada paraula; després llegeix l’exemple en veu alta."),
        gtip=("Gramàtica: ", "llegeix com funciona, tapa la columna de respostes i completa les frases; comprova-ho i llegeix el perquè."),
        vocab="1 · Vocabulari", grammar="2 · Gramàtica",
        gintro="El Learning Path 1 és un escalfament i el 6 és el producte final: no tenen gramàtica nova.",
        cols=("Paraula o expressió", "En català", "Exemple"), how="Com funciona",
        practice="Pràctica del Learning Path, amb respostes", pcols=("Frase", "Resposta", "Per què"),
    ),
}
GRAMMAR_PAGE = {  # title of the grammar page block (LP4) per course and route
    ("4eso", "insecure"): "Present Simple", ("4eso", "safe"): "Present Simple and phrasal verbs",
    ("4eso", "daring"): "Present Simple and phrasal verbs", ("2eso", "insecure"): "Present Simple",
    ("2eso", "safe"): "Present Simple", ("2eso", "daring"): "Present Simple",
}


def rgb(h):
    return RGBColor.from_string(h)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill)
    tcPr.append(sh)


def borders(table, color=BORDER):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + side)
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def margins(cell, top=50, bottom=50):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for side, v in (("top", top), ("bottom", bottom), ("left", 90), ("right", 90)):
        e = OxmlElement("w:" + side); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tcPr.append(m)


def keep_with_next(p):
    p.paragraph_format.keep_with_next = True


def run(p, text, bold=False, italic=False, size=None, color=None, underline=False, shading=None):
    r = p.add_run(text)
    r.bold, r.italic, r.underline = bold, italic, underline
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = rgb(color)
    if shading:
        rPr = r._r.get_or_add_rPr()
        sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), shading)
        rPr.append(sh)
    return r


class Html(HTMLParser):
    """Turns the explanation HTML of the site into paragraphs of runs: ('text'|'model', [(text, fmt)])."""

    def __init__(self):
        super().__init__()
        self.paras, self.fmt, self.kind = [("text", [])], [], "text"

    def handle_starttag(self, tag, attrs):
        cls = dict(attrs).get("class", "")
        if tag == "div":
            self.kind = "model"; self.paras.append(("model", []))
        elif tag == "br":
            self.paras.append((self.kind, []))
        elif tag in ("strong", "b"):
            self.fmt.append("bold")
        elif tag == "em":
            self.fmt.append("italic")
        elif tag == "span":
            self.fmt.append({"verb": "bold", "particle": "shade", "object": "underline"}.get(cls, ""))

    def handle_endtag(self, tag):
        if tag == "div":
            self.kind = "text"; self.paras.append(("text", []))
        elif tag in ("strong", "b", "em", "span") and self.fmt:
            self.fmt.pop()

    def handle_data(self, data):
        self.paras[-1][1].append((data, set(self.fmt)))


def html_paras(html):
    h = Html(); h.feed(html)
    out = []
    for kind, runs in h.paras:
        runs = [(t, f) for t, f in runs if t]
        if "".join(t for t, _ in runs).strip():
            out.append((kind, runs))
    return out


class Builder:
    def __init__(self, data):
        self.d, self.course = data, data["course"]
        self.doc = Document()
        s = self.doc.styles["Normal"]
        s.font.name, s.font.size = "Calibri", Pt(10.5)
        s.font.color.rgb = rgb(TEXT)
        s.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        s.paragraph_format.space_after = Pt(4)
        for name, size in (("Heading 1", 17), ("Heading 2", 13.5), ("Heading 3", 12)):
            h = self.doc.styles[name]
            h.font.name, h.font.size, h.font.bold = "Calibri", Pt(size), True
            h.font.color.rgb = rgb(ACCENT)
            rf = h.element.rPr.rFonts
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(a), "Calibri")
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            h.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
            h.paragraph_format.space_after = Pt(4)
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(1.8)
        sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.8)
        self.width = sec.page_width - sec.left_margin - sec.right_margin
        self.footer(sec)
        self.first = True

    def footer(self, sec):
        p = sec.footer.paragraphs[0]
        p.style = self.doc.styles["Normal"]
        p.paragraph_format.tab_stops.add_tab_stop(self.width, WD_TAB_ALIGNMENT.RIGHT)
        run(p, COPYRIGHT + "\t", size=8, color=GREY)
        run(p, "%s · %s · Vocabulary and grammar · page " % (LS, COURSE[self.course]), size=8, color=GREY)
        r = p.add_run(); r.font.size = Pt(8); r.font.color.rgb = rgb(GREY)
        for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if kind:
                e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind); r._r.append(e)
            else:
                e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = text; r._r.append(e)

    def para(self, text="", **kw):
        p = self.doc.add_paragraph()
        if text: run(p, text, **kw)
        return p

    def new_page(self):
        if not self.first:
            self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        self.first = False

    def box(self, title, lines):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders(t, BORDER)
        c = t.cell(0, 0); shade(c, LIGHT); margins(c, 110, 110)
        p = c.paragraphs[0]; run(p, title, bold=True, color=ACCENT)
        for line in lines:
            p = c.add_paragraph(); p.paragraph_format.space_after = Pt(2)
            for text, bold in line:
                run(p, text, bold=bold)
        self.para()

    def table(self, header, rows, widths, group_rows=(), bold_col=None, italic_col=None):
        t = self.doc.add_table(rows=1, cols=len(header))
        borders(t)
        t.autofit = False
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]; shade(c, ACCENT); margins(c)
            run(c.paragraphs[0], h, bold=True, color="FFFFFF", size=9.5)
        tr = t.rows[0]._tr.get_or_add_trPr()
        e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); tr.append(e)
        for i, row in enumerate(rows):
            cells = t.add_row().cells
            if i in group_rows:
                m = cells[0].merge(cells[-1]); shade(m, LIGHT); margins(m)
                run(m.paragraphs[0], row[0], bold=True, color=ACCENT, size=9.5)
                continue
            for j, val in enumerate(row):
                margins(cells[j])
                run(cells[j].paragraphs[0], val, bold=(j == bold_col), italic=(j == italic_col), size=9.5)
        for row in t.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for j, w in enumerate(widths):
                row.cells[j].width = int(self.width * w)
                row.cells[j].paragraphs[0].paragraph_format.space_after = Pt(0)
        self.para().paragraph_format.space_after = Pt(2)

    # ---- one route part -------------------------------------------------------------------------------------------
    def visible(self, item, route):
        if item.get("only"):
            return item["route"] == route
        return ROUTES.index(item["route"]) <= ROUTES.index(route)

    def tag(self, item, route, lang):
        if lang == "ca" or item.get("only") or item["route"] == "insecure":
            return ""
        return " (%s)" % NAME[item["route"]]

    def part(self, route, heading=None):
        lang = "ca" if route == "insecure" else "en"
        s = T[lang]
        self.new_page()
        self.para("%s · %s" % (COURSE[self.course], LS), bold=True, size=9, color=ACCENT)
        p = self.para(heading or s["route"].format(r=NAME[route]), bold=True, size=20, color=ACCENT)
        p.paragraph_format.space_after = Pt(6)
        self.para(s["intro"].format(r=NAME[route]))
        self.box(s["box"], [[(s["order"], False)], [(s["vtip"][0], True), (s["vtip"][1], False)],
                            [(s["gtip"][0], True), (s["gtip"][1], False)]])
        self.doc.add_heading(s["vocab"], 1)
        for lp in self.d["lps"]:
            groups = [g for g in lp["vocab"] if self.visible(g, route)]
            if not groups:
                continue
            keep_with_next(self.doc.add_heading("Learning Path %d · %s" % (lp["n"], lp["title"]), 2))
            rows, group_rows = [], []
            for g in groups:
                name = g.get("nameCa") if lang == "ca" and g.get("nameCa") else g["name"]
                group_rows.append(len(rows))
                rows.append(["%s%s · %d" % (name, self.tag(g, route, lang), len(g["items"]))])
                rows += [list(x) for x in g["items"]]
            self.table(s["cols"], rows, (0.34, 0.30, 0.36), group_rows, bold_col=0, italic_col=2)
        self.doc.add_heading(s["grammar"], 1)
        self.para(s["gintro"], color=GREY)
        for lp in self.d["lps"]:
            blocks = [b for b in lp["grammar"] if self.visible(b, route)]
            if not blocks:
                continue
            keep_with_next(self.doc.add_heading("Learning Path %d · %s" % (lp["n"], lp["title"]), 2))
            for b in blocks:
                title = b["title"] or GRAMMAR_PAGE[(self.course, route)]
                keep_with_next(self.doc.add_heading(title + self.tag(b, route, lang), 3))
                if b.get("how") or b.get("lessons"):
                    keep_with_next(self.para(s["how"], bold=True, size=10, color=ACCENT))
                for line in b.get("how") or []:
                    self.bullet_html(line)
                for title_l, html in b.get("lessons") or []:
                    p = self.para(title_l, bold=True, size=10.5)
                    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(1)
                    keep_with_next(p)
                    for kind, runs in html_paras(html):
                        q = self.doc.add_paragraph()
                        q.paragraph_format.space_after = Pt(2)
                        if kind == "model":
                            q.paragraph_format.left_indent = Cm(0.8)
                        for text, fmt in runs:
                            run(q, text, bold="bold" in fmt, italic=(kind == "model") or "italic" in fmt,
                                underline="underline" in fmt, shading=BORDER if "shade" in fmt else None)
                if b.get("practice"):
                    p = self.para(s["practice"], bold=True, size=10, color=ACCENT)
                    p.paragraph_format.space_before = Pt(6); keep_with_next(p)
                    if any(x[2] for x in b["practice"]):
                        self.table(s["pcols"], b["practice"], (0.52, 0.22, 0.26), bold_col=1)
                    else:
                        self.table(s["pcols"][:2], [x[:2] for x in b["practice"]], (0.66, 0.34), bold_col=1)

    def bullet_html(self, text):
        p = self.doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        if ": " in text and len(text.split(": ", 1)[0]) < 40:
            head, rest = text.split(": ", 1)
            run(p, head + ": ", bold=True); run(p, rest)
        else:
            run(p, text)

    # ---- teacher's Word -------------------------------------------------------------------------------------------
    def teacher(self):
        self.first = False
        self.para("%s · %s" % (COURSE[self.course], LS), bold=True, size=9, color=ACCENT)
        self.para("Grammar and vocabulary by route", bold=True, size=22, color=ACCENT)
        self.para("Everything the Learning Paths teach, with the explanations as students see them", italic=True, color=GREY)
        self.para("Built from the live site data (the Learning Path pages of LS1) on %s. If a Learning Path changes, this "
                  "document can be rebuilt from it (tools/study-summary in the repo)." % datetime.date.today().strftime("%-d %b %Y"))
        self.box("How to use this document", [
            [("This page is for you: a map of what each route learns in each Learning Path. Then there is one part per route, each "
              "starting on its own page, ready to print or post on Classroom: Part 1 Insecure (in Catalan, as on the site), Part 2 "
              "Safe and Part 3 Daring (in English).", False)],
            [("The vocabulary builds up, as on the site: Safe students see the Insecure words plus the Safe ones, and Daring students "
              "see all three levels; a tag such as (Safe) or (Daring) shows where a group starts. The grammar page, the Useful "
              "expressions / Same day practice and the connectors have their own version per route, as on the site.", False)],
            [("Beginners (Starter): ", True), ("LS1 has no Beginners version on the site, so Beginners students see the Insecure "
              "pages; give them Part 1 (Insecure).", False)],
            [("Students can download the same parts as PDFs (only their route).", False)],
        ])
        self.doc.add_heading("What each route learns", 1)
        rows = []
        for lp in self.d["lps"]:
            cells = []
            for r in ROUTES:
                lines = []
                words = [g for g in lp["vocab"] if g["route"] == r]
                if words:
                    lines.append("Words: " + ", ".join("%s (%d)" % (g["name"], len(g["items"])) for g in words))
                gr = [b for b in lp["grammar"] if b["route"] == r]
                if gr:
                    parts = []
                    for b in gr:
                        if b.get("lessons"):
                            parts.append(" · ".join(re.sub(r"^\d+ · ", "", t) for t, _ in b["lessons"]))
                        else:
                            parts.append(b["title"])
                    lines.append("Grammar: " + " · ".join(parts))
                cells.append("\n".join(lines) or ("None new" if lp.get("none") and r == "insecure" else "—"))
            rows.append(["%d · %s" % (lp["n"], lp["title"])] + cells)
        self.table(("Learning Path", "Insecure (every route)", "Safe", "Daring"), rows, (0.19, 0.29, 0.26, 0.26), bold_col=0)
        self.para("The numbers in brackets are the number of words or expressions in each group. Words build up by route; "
                  "grammar items marked per route are that route's own version. Not included: listening and reading questions, "
                  "the vocabulary games, open writing and speaking tasks and self-checks.", size=9, color=GREY)
        for i, r in enumerate(ROUTES):
            lang = "ca" if r == "insecure" else "en"
            label = ("Part %d · " % (i + 1)) + T[lang]["route"].format(r=NAME[r])
            self.part(r, heading=label)


def main():
    out, data = sys.argv[1], json.load(open(sys.argv[2], encoding="utf-8"))
    b = Builder(data)
    if len(sys.argv) > 3:
        b.part(sys.argv[3])
    else:
        b.teacher()
    b.doc.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
